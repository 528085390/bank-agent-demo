"""表格：CSV（标准库）与 XLSX（openpyxl）。

两者都渲染成 TSV 风格的 ``text``（``# sheet:`` + 制表符）：下游的字符串解析器按
制表符分列，原样交出逗号文本会让它一列都切不开。结构化数据同时以二维数组放进
``metadata["sheets"]``，表头与行列数供下游取用。

这里**不做行数截断**：截断一张对账单等于给下游一份不完整的账，宁可让上游的
压缩包尺寸上限去兜底（见 ``ArchiveLimits.max_total_size``）。
"""

from __future__ import annotations

import csv
import io
from datetime import date, datetime, time
from typing import Any

from app.billing import detect
from app.billing.contracts import AttachmentResult
from app.billing.detect import decode_text
from app.billing.errors import missing_dependency, unsupported
from app.billing.handlers.registry import HandleContext, result_for


class SpreadsheetHandler:
    kind = "spreadsheet"

    def handle(self, data: bytes, ctx: HandleContext) -> AttachmentResult:
        if ctx.detected_type == detect.CSV:
            return self._handle_csv(data, ctx)
        return self._handle_xlsx(data, ctx)

    def _handle_csv(self, data: bytes, ctx: HandleContext) -> AttachmentResult:
        decoded = decode_text(data, declared=detect.CSV)
        if decoded is None:
            raise unsupported("无法按 UTF-8 或 GB18030 解码 CSV")
        text, encoding = decoded
        rows = list(csv.reader(io.StringIO(text)))
        # CSV 与 XLSX 走同一条渲染路径：下游的解析器（parse_cmb_statement）按制表符分列，
        # 原样交出逗号文本会让它一列都切不开，整份账单解析为空。metadata 保留原有的
        # encoding / rows / columns / header 四个键，只额外给出渲染用的 sheet 结构。
        sheet = _csv_sheet(ctx.filename, rows)
        return result_for(
            data,
            ctx,
            text=_render_sheets_as_text([sheet]),
            metadata={
                "encoding": encoding,
                "rows": sheet["rows"],
                "columns": sheet["columns"],
                "header": sheet["header"],
                "sheets": [sheet],
            },
        )

    def _handle_xlsx(self, data: bytes, ctx: HandleContext) -> AttachmentResult:
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise missing_dependency("openpyxl", "解析 XLSX") from exc

        try:
            workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception as exc:
            raise unsupported("XLSX 已损坏或不是有效的工作簿") from exc

        try:
            sheets = [_read_sheet(sheet) for sheet in workbook.worksheets]
        finally:
            workbook.close()

        return result_for(
            data,
            ctx,
            text=_render_sheets_as_text(sheets),
            metadata={"sheets": sheets},
        )


def _csv_sheet(name: str, rows: list[list[str]]) -> dict[str, Any]:
    """CSV 的行拼成与 ``_read_sheet`` 同形的 sheet，好复用同一个渲染器。

    表名取文件名去掉扩展名：CSV 没有工作表概念，但 ``# sheet:`` 需要一个标识。
    """

    return {
        "name": name.rsplit(".", 1)[0] if name else "csv",
        "rows": len(rows),
        "columns": max((len(row) for row in rows), default=0),
        "header": rows[0] if rows else [],
        "values": rows[1:] if rows else [],
    }


def _read_sheet(sheet: Any) -> dict[str, Any]:
    values = [[_cell(value) for value in row] for row in sheet.iter_rows(values_only=True)]
    return {
        "name": sheet.title,
        "rows": len(values),
        "columns": max((len(row) for row in values), default=0),
        "header": values[0] if values else [],
        "values": values[1:] if values else [],
    }


def _cell(value: Any) -> Any:
    """单元格值：时间类型转 ISO 字符串，其余保留原样。"""

    if value is None or isinstance(value, (bool, int, float, str)):
        return value
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    return str(value)


def _render_sheets_as_text(sheets: list[dict[str, Any]]) -> str:
    """把各工作表渲染为 TSV 风格文本，供下游的字符串解析器消费。

    单元格一律经 ``_cell`` 归一（见 ``_cell_to_text``），因此日期在 text 与
    ``metadata["sheets"]`` 里是同一个 ISO 表示，不会出现两种格式。不截断、不改写
    单元格内容。``metadata["sheets"]`` 仍是结构化数据的权威副本，这里只多给一份
    可解析的文本视图，供 ``parse_cmb_statement(text)`` 这类字符串解析器使用。
    """
    blocks = []
    for sheet in sheets:
        rows = ([sheet["header"]] if sheet["header"] else []) + sheet["values"]
        if not rows:
            continue
        lines = ["\t".join(_cell_to_text(value) for value in row) for row in rows]
        blocks.append(f"# sheet: {sheet['name']}\n" + "\n".join(lines))
    return "\n\n".join(blocks)


def _cell_to_text(value: Any) -> str:
    """单元格渲染用：复用 ``_cell`` 的归一结果，None 记空串。"""

    normalized = _cell(value)
    return "" if normalized is None else str(normalized)
