"""DOCX：用 python-docx 提取段落与表格文本。

依赖按需导入。方案 §3.1 把 DOCX 排在批次 2，所以 ``python-docx`` 目前不在
``backend/pyproject.toml`` 里：装上即可用（``uv add python-docx``），没装时
只有这一条路报 ATTACHMENT_UNSUPPORTED 并给出安装命令。
"""

from __future__ import annotations

import io

from app.billing.contracts import AttachmentResult
from app.billing.errors import missing_dependency, unsupported
from app.billing.handlers.registry import HandleContext, result_for


class DocxHandler:
    kind = "document"

    def handle(self, data: bytes, ctx: HandleContext) -> AttachmentResult:
        try:
            import docx
        except ImportError as exc:
            raise missing_dependency("python-docx", "解析 DOCX") from exc

        try:
            document = docx.Document(io.BytesIO(data))
        except Exception as exc:
            raise unsupported("DOCX 已损坏或无法解析") from exc

        paragraphs = [paragraph.text for paragraph in document.paragraphs]
        tables = [
            [[cell.text for cell in row.cells] for row in table.rows] for table in document.tables
        ]

        lines = [line for line in paragraphs if line.strip()]
        for table in tables:
            lines.extend("\t".join(cells) for cells in table)
        text = "\n".join(lines)

        return result_for(
            data,
            ctx,
            text=text or None,
            metadata={
                "paragraphs": len(paragraphs),
                "tables": len(tables),
                "table_rows": sum(len(table) for table in tables),
            },
        )
